import os
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

os.makedirs("sample_resumes", exist_ok=True)

# 1. Generate DOCX Sample Resume
doc = Document()
doc.add_heading("Alex Morgan", level=0)
doc.add_paragraph("Seattle, WA | (206) 555-0199 | alex.morgan@email.com | linkedin.com/in/alexmorgan | github.com/alexmorgan")

doc.add_heading("Professional Summary", level=1)
doc.add_paragraph("Full Stack Software Engineer with 5 years of experience building web applications and cloud architectures. Proficient in React, TypeScript, Python, and AWS.")

doc.add_heading("Work Experience", level=1)
doc.add_paragraph("Senior Software Engineer | Apex Cloud Solutions | 2022 - Present")
doc.add_paragraph("• Architected scalable microservices using FastAPI and Python, reducing API response latency by 42% for 120,000 active users.", style='List Bullet')
doc.add_paragraph("• Spearheaded migration from legacy monolithic MySQL database to PostgreSQL cluster, cutting downtime by 99.9%.", style='List Bullet')
doc.add_paragraph("• Worked on React dashboard components and fixed various frontend UI bugs.", style='List Bullet')

doc.add_paragraph("Software Engineer | DevMatrix Systems | 2019 - 2022")
doc.add_paragraph("• Developed RESTful endpoints using Node.js and Express, supporting 4 enterprise partner integrations.", style='List Bullet')
doc.add_paragraph("• Responsible for writing unit test suites and maintaining 85% test coverage in CI/CD pipeline.", style='List Bullet')
doc.add_paragraph("• Helped with cloud server monitoring and dockerized dev environments.", style='List Bullet')

doc.add_heading("Education", level=1)
doc.add_paragraph("Bachelor of Science in Computer Science | University of Washington | 2015 - 2019")

doc.add_heading("Technical Skills", level=1)
doc.add_paragraph("Languages: Python, TypeScript, JavaScript, SQL, Go\nFrameworks: FastAPI, React, Next.js, Node.js, Express\nCloud & Tools: AWS, Docker, Kubernetes, Git, PostgreSQL, Redis, CI/CD")

doc.add_heading("Projects", level=1)
doc.add_paragraph("Distributed Task Queue Engine\n• Built an open-source background task processor in Go and Redis with 500+ GitHub stars.")

doc.add_heading("Certifications", level=1)
doc.add_paragraph("AWS Certified Solutions Architect - Associate")

docx_path = "sample_resumes/sample_software_engineer.docx"
doc.save(docx_path)
print(f"Generated DOCX sample: {docx_path}")

# 2. Generate PDF Sample Resume
pdf_path = "sample_resumes/sample_product_manager.pdf"
pdf = SimpleDocTemplate(pdf_path, pagesize=letter, leftMargin=0.5*inch, rightMargin=0.5*inch, topMargin=0.5*inch, bottomMargin=0.5*inch)
styles = getSampleStyleSheet()

story = []
title_style = ParagraphStyle('Title', parent=styles['Heading1'], fontSize=18, leading=22, textColor=colors.HexColor('#0f172a'))
h2_style = ParagraphStyle('H2', parent=styles['Heading2'], fontSize=12, leading=16, textColor=colors.HexColor('#2563eb'), spaceBefore=8, spaceAfter=4)
body_style = ParagraphStyle('Body', parent=styles['Normal'], fontSize=9, leading=13, textColor=colors.HexColor('#334155'))

story.append(Paragraph("<b>Jordan Lee</b>", title_style))
story.append(Paragraph("New York, NY | (212) 555-7821 | jordan.lee@email.com | linkedin.com/in/jordanlee", body_style))
story.append(Spacer(1, 6))

story.append(Paragraph("<b>PROFESSIONAL SUMMARY</b>", h2_style))
story.append(Paragraph("Strategic Lead Product Manager with 6+ years driving product vision, roadmap execution, and high-growth B2B SaaS initiatives. Adept in agile methodologies, customer discovery, data analytics, and cross-functional team leadership.", body_style))

story.append(Paragraph("<b>WORK EXPERIENCE</b>", h2_style))
story.append(Paragraph("<b>Lead Product Manager</b> | FinScale Technologies | 2021 - Present", body_style))
story.append(Paragraph("• Launched enterprise billing product line generating $2.4M ARR in its first 12 months with 18 enterprise logos.", body_style))
story.append(Paragraph("• Increased user activation rate by 24% by redesigning client onboarding flows based on mixed-methods user research.", body_style))
story.append(Paragraph("• Managed sprint priorities across 3 engineering squads totaling 18 engineers and 2 product designers.", body_style))

story.append(Paragraph("<b>Associate Product Manager</b> | VenturePulse Apps | 2018 - 2021", body_style))
story.append(Paragraph("• Handled customer user interviews and wrote user stories for the mobile app team.", body_style))
story.append(Paragraph("• Analyzed user retention cohorts using Mixpanel and SQL to identify key drop-off triggers.", body_style))

story.append(Paragraph("<b>EDUCATION</b>", h2_style))
story.append(Paragraph("B.S. in Information Systems & Business Analytics | New York University | 2014 - 2018", body_style))

story.append(Paragraph("<b>CORE COMPETENCIES & SKILLS</b>", h2_style))
story.append(Paragraph("Product Strategy, Roadmapping, Agile/Scrum, User Discovery, A/B Testing, Wireframing, SQL, Jira, Mixpanel, Market Research, Stakeholder Alignment", body_style))

pdf.build(story)
print(f"Generated PDF sample: {pdf_path}")
